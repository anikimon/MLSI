"""Excel intake and weighted descriptive reports for research projects."""

import io
import json
import math
import os
import re
import zipfile
from collections import defaultdict
from html import escape
from pathlib import Path
from xml.etree.ElementTree import ParseError

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.enums import TA_JUSTIFY
from reportlab.lib.units import cm
from reportlab.graphics.shapes import Drawing, Rect
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


def cell_text(value):
    return "" if value is None else str(value).strip()


def parse_weight(value):
    try:
        weight = float(str(value).replace(",", "."))
    except (TypeError, ValueError):
        raise ValueError("Вес должен быть числом больше нуля и не больше 1000")
    if not math.isfinite(weight) or not 0 < weight <= 1000:
        raise ValueError("Вес должен быть числом больше нуля и не больше 1000")
    return weight


def read_excel(content, questions):
    if not zipfile.is_zipfile(io.BytesIO(content)):
        raise ValueError("Поддерживаются только файлы Excel .xlsx")
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        if sum(item.file_size for item in archive.infolist()) > 50 * 1024 * 1024:
            raise ValueError("Распакованный Excel-файл слишком большой")
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True, keep_links=False)
        try:
            rows = workbook.worksheets[0].iter_rows(values_only=True)
            headers = [cell_text(value) for value in next(rows)]
            if not headers or any(not value for value in headers) or len(headers) > 42:
                raise ValueError("Первая строка должна содержать названия колонок (не более 42)")
            normalized = [value.casefold() for value in headers]
            if len(set(normalized)) != len(headers):
                raise ValueError("Названия колонок не должны повторяться")
            code_col = normalized.index("code") if "code" in normalized else None
            weight_col = normalized.index("weight") if "weight" in normalized else None
            question_cols = [i for i in range(len(headers)) if i not in (code_col, weight_col)]
            if not question_cols:
                raise ValueError("Добавьте в Excel хотя бы один столбец с ответами")
            if questions:
                columns = {}
                for question in questions:
                    possible = [i for i in question_cols if headers[i] in (question["id"], question["label"])]
                    if len(possible) != 1:
                        raise ValueError(f"Нет однозначной колонки для вопроса: {question['label']}")
                    columns[question["id"]] = possible[0]
                if set(columns.values()) != set(question_cols):
                    raise ValueError("В файле есть лишние колонки с ответами")
            else:
                if len(question_cols) > 40:
                    raise ValueError("Допускается не более 40 вопросов")
                columns = {f"q{i + 1}": column for i, column in enumerate(question_cols)}

            records = []
            for row_number, row in enumerate(rows, start=2):
                if row_number > 10001:
                    raise ValueError("За один раз можно загрузить не более 10 000 анкет")
                if not any(value is not None and str(value).strip() for value in row):
                    continue
                if len(row) > len(headers) and any(cell_text(v) for v in row[len(headers):]):
                    raise ValueError(f"Строка {row_number}: данные без заголовка")
                values = list(row) + [None] * (len(headers) - len(row))
                try:
                    weight = parse_weight(values[weight_col]) if weight_col is not None else 1.0
                except ValueError as exc:
                    raise ValueError(f"Строка {row_number}: {exc}")
                record = {
                    "row_number": row_number,
                    "code": cell_text(values[code_col]) if code_col is not None else "",
                    "weight": weight,
                    "answers": {key: cell_text(values[column]) for key, column in columns.items()},
                }
                records.append(record)
            if not records:
                raise ValueError("В Excel-файле нет заполненных анкет")
            if not questions:
                questions = []
                for key, column in columns.items():
                    values = [record["answers"][key] for record in records if record["answers"][key]]
                    distinct = list(dict.fromkeys(values))
                    numeric = bool(values) and all(re.fullmatch(r"-?\d+(?:[.,]\d+)?", value) for value in values)
                    multi = any(";" in value for value in values)
                    choices = list(dict.fromkeys(v.strip() for value in values for v in value.split(";") if v.strip()))
                    if numeric:
                        kind, options = "number", []
                    elif multi and 2 <= len(choices) <= 20:
                        kind, options = "multiple", choices
                    elif 2 <= len(distinct) <= 20:
                        kind, options = "single", distinct
                    else:
                        kind, options = "text", []
                    questions.append({"id": key, "label": headers[column], "type": kind, "required": False, "options": options})
            for record in records:
                for question in questions:
                    if question["type"] == "multiple":
                        record["answers"][question["id"]] = [part.strip() for part in record["answers"][question["id"]].split(";") if part.strip()]
            return questions, records
        finally:
            workbook.close()
    except (IndexError, StopIteration, KeyError, zipfile.BadZipFile, InvalidFileException, ParseError, OSError) as exc:
        raise ValueError("Не удалось прочитать Excel-файл") from exc


def build_quotas(db, study_id, questions):
    configs = db.execute("SELECT question_id, rules FROM study_quotas WHERE study_id = ? ORDER BY question_id", (study_id,)).fetchall()
    if not configs:
        return []
    answers = [json.loads(row["answers"]) for row in db.execute("SELECT answers FROM responses WHERE study_id = ?", (study_id,))]
    by_id = {question["id"]: question for question in questions}
    result = []
    for config in configs:
        question = by_id[config["question_id"]]
        rows = []
        for rule in json.loads(config["rules"]):
            if question["type"] == "number":
                label = f"{rule['min']:g}–{rule['max']:g}"
                count = sum(1 for item in answers if item.get(config["question_id"]) and
                            rule["min"] <= float(str(item[config["question_id"]]).replace(",", ".")) < rule["max"])
            else:
                label = rule["option"]
                count = sum(rule["option"] in item.get(config["question_id"], []) if question["type"] == "multiple"
                            else item.get(config["question_id"]) == rule["option"] for item in answers)
            target = rule["target"]
            rows.append({**rule, "label": label, "count": count, "remaining": max(target - count, 0),
                         "percent": round(100 * count / target, 1) if target else None})
        result.append({"question_id": question["id"], "label": question["label"], "type": question["type"], "rows": rows})
    return result


def build_member_quotas(db, study_id):
    rows = db.execute("""SELECT q.id, q.user_id, u.name AS user_name, q.label, q.conditions, q.target FROM member_quotas q
        JOIN users u ON u.id = q.user_id WHERE q.study_id = ? ORDER BY q.user_id, q.id""", (study_id,)).fetchall()
    if not rows:
        return []
    answers = defaultdict(list)
    for row in db.execute("SELECT interviewer_id, answers FROM responses WHERE study_id = ? AND source = 'online'", (study_id,)):
        answers[row["interviewer_id"]].append(json.loads(row["answers"]))
    result = []
    for row in rows:
        conditions = json.loads(row["conditions"])
        count = sum(all((condition["option"] in item.get(condition["question_id"], []) if isinstance(item.get(condition["question_id"]), list)
                         else item.get(condition["question_id"]) == condition["option"])
                        if "option" in condition else (
                            bool(item.get(condition["question_id"])) and
                            (condition.get("min") is None or float(item[condition["question_id"]].replace(",", ".")) >= condition["min"]) and
                            (condition.get("max") is None or float(item[condition["question_id"]].replace(",", ".")) < condition["max"]))
                        for condition in conditions) for item in answers[row["user_id"]])
        result.append({"id": row["id"], "user_id": row["user_id"], "user_name": row["user_name"], "label": row["label"], "conditions": conditions,
                       "target": row["target"], "count": count, "remaining": max(0, row["target"] - count),
                       "percent": round(100 * count / row["target"], 1) if row["target"] else None})
    return result


def build_association(questions, prepared, x_id, y_id):
    by_id = {question["id"]: question for question in questions}
    x, y = by_id[x_id], by_id[y_id]
    pairs = []
    for answers, weight in prepared:
        a, b = answers.get(x_id), answers.get(y_id)
        if a in (None, "") or b in (None, ""):
            continue
        a = float(str(a).replace(",", ".")) if x["type"] == "number" else a
        b = float(str(b).replace(",", ".")) if y["type"] == "number" else b
        if (isinstance(a, float) and not math.isfinite(a)) or (isinstance(b, float) and not math.isfinite(b)):
            continue
        pairs.append((a, b, weight))
    total = sum(weight for _, _, weight in pairs)
    result = {"x_id": x_id, "y_id": y_id, "x_label": x["label"], "y_label": y["label"],
              "count": len(pairs), "weighted_base": round(total, 2), "value": None, "details": []}
    if x["type"] == y["type"] == "number":
        result["method"] = "Пирсон r"
        if len(pairs) >= 3:
            ax = sum(a * w for a, _, w in pairs) / total
            ay = sum(b * w for _, b, w in pairs) / total
            vx = sum(w * (a - ax) ** 2 for a, _, w in pairs)
            vy = sum(w * (b - ay) ** 2 for _, b, w in pairs)
            if vx > 0 and vy > 0:
                result["value"] = round(max(-1, min(1, sum(w * (a - ax) * (b - ay) for a, b, w in pairs) / math.sqrt(vx * vy))), 3)
    elif x["type"] == y["type"] == "single":
        result["method"] = "V Крамера"
        matrix = [[sum(w for a, b, w in pairs if a == xa and b == yb) for yb in y["options"]] for xa in x["options"]]
        result["details"] = [{"label": label, "values": [round(value, 2) for value in row]}
                             for label, row in zip(x["options"], matrix)]
        result["columns"] = y["options"]
        row_sums = [sum(row) for row in matrix]
        col_sums = [sum(row[i] for row in matrix) for i in range(len(y["options"]))]
        active_rows = sum(value > 0 for value in row_sums)
        active_cols = sum(value > 0 for value in col_sums)
        if len(pairs) >= 3 and active_rows > 1 and active_cols > 1:
            chi = sum((value - row_sums[i] * col_sums[j] / total) ** 2 / (row_sums[i] * col_sums[j] / total)
                      for i, row in enumerate(matrix) for j, value in enumerate(row) if row_sums[i] and col_sums[j])
            result["value"] = round(min(1, math.sqrt(chi / (total * min(active_rows - 1, active_cols - 1)))), 3)
    else:
        result["method"] = "Отношение корреляции η"
        numeric_first = x["type"] == "number"
        group_question = y if numeric_first else x
        groups = [(option, [(a if numeric_first else b, w) for a, b, w in pairs if (b if numeric_first else a) == option])
                  for option in group_question["options"]]
        mean = sum((a if numeric_first else b) * w for a, b, w in pairs) / total if total else 0
        variance = sum(w * ((a if numeric_first else b) - mean) ** 2 for a, b, w in pairs)
        between = 0
        for option, values in groups:
            base = sum(w for _, w in values)
            group_mean = sum(v * w for v, w in values) / base if base else None
            result["details"].append({"label": option, "count": len(values), "weighted_base": round(base, 2),
                                      "mean": round(group_mean, 2) if group_mean is not None else None})
            if base:
                between += base * (group_mean - mean) ** 2
        if len(pairs) >= 3 and sum(bool(values) for _, values in groups) > 1 and variance > 0:
            result["value"] = round(min(1, math.sqrt(between / variance)), 3)
    if result["value"] is None:
        result["analysis"] = "Связь не рассчитана: нужно не менее трёх пар ответов и различающиеся значения обеих переменных."
    else:
        result["analysis"] = (f"{result['method']} = {result['value']:.3f} по {len(pairs)} полным парам ответов. "
                              "Показатель описывает связь в выборке, но не доказывает причинно-следственную связь.")
    return result


def build_report(db, study_id):
    study = db.execute("SELECT title, stage, goal, tasks FROM studies WHERE id = ?", (study_id,)).fetchone()
    row = db.execute("SELECT questions FROM questionnaires WHERE study_id = ?", (study_id,)).fetchone()
    questions = json.loads(row["questions"]) if row else []
    rule = db.execute("SELECT question_id, coefficients FROM weighting WHERE study_id = ?", (study_id,)).fetchone()
    coefficients = json.loads(rule["coefficients"]) if rule else {}
    responses = db.execute("SELECT answers, weight, source FROM responses WHERE study_id = ?", (study_id,)).fetchall()
    prepared = [(answers, response["weight"] * coefficients.get(answers.get(rule["question_id"]), 1) if rule else response["weight"])
                for response in responses for answers in [json.loads(response["answers"])]]
    association_row = db.execute("SELECT x_id, y_id FROM study_associations WHERE study_id = ?", (study_id,)).fetchone()
    association = build_association(questions, prepared, association_row["x_id"], association_row["y_id"]) if association_row else None
    quotas = build_quotas(db, study_id, questions)
    count_online = sum(response["source"] == "online" for response in responses)
    count_excel = sum(response["source"] == "excel" for response in responses)
    count_public = sum(response["source"] == "public" for response in responses)
    result = {
        "title": study["title"], "stage": study["stage"], "goal": study["goal"], "tasks": study["tasks"],
        "count": len(responses),
        "online_count": count_online, "excel_count": count_excel, "public_count": count_public,
        "total_weight": round(sum(weight for _, weight in prepared), 2), "questions": [],
        "refusal_count": db.execute("SELECT COUNT(*) FROM refusals WHERE study_id = ?", (study_id,)).fetchone()[0],
        "weighting": {"question_id": rule["question_id"], "coefficients": coefficients} if rule else None,
        "weighting_questions": [{"id": q["id"], "label": q["label"], "options": q["options"]}
                                for q in questions if q["type"] == "single"],
        "association_questions": [{"id": q["id"], "label": q["label"], "type": q["type"]}
                                  for q in questions if q["type"] in ("single", "number")],
        "association": association, "quotas": quotas, "member_quotas": build_member_quotas(db, study_id),
    }
    for question in questions:
        key = question["id"]
        totals = defaultdict(lambda: [0, 0.0])
        base = 0.0
        answered = 0
        numeric_sum = 0.0
        numbers = []
        for answers, weight in prepared:
            answer = answers.get(key, "")
            if not answer:
                continue
            base += weight
            answered += 1
            if question["type"] == "number":
                number = float(str(answer).replace(",", "."))
                numeric_sum += number * weight
                numbers.append((number, weight))
            elif question["type"] == "multiple":
                for choice in set(answer):
                    totals[choice][0] += 1
                    totals[choice][1] += weight
            else:
                totals[answer][0] += 1
                totals[answer][1] += weight
        if numbers:
            low = min(value for value, _ in numbers)
            high = max(value for value, _ in numbers)
            width = (high - low) / 8
            categories = [f"{low:.12g}"] if not width else [f"{low + i * width:.12g}–{low + (i + 1) * width:.12g}" for i in range(8)]
            # Keep bin indices separate: rounded labels can coincide for narrow ranges.
            bins = [[0, 0.0] for _ in categories]
            for value, weight in numbers:
                bucket = min(int((value - low) / width), 7) if width else 0
                bins[bucket][0] += 1
                bins[bucket][1] += weight
            rows = [{"label": label, "count": count, "weighted_count": round(weighted, 2),
                     "percent": round(100 * weighted / base, 1)} for label, (count, weighted) in zip(categories, bins)]
        else:
            categories = question["options"] if question["type"] in ("single", "multiple") else sorted(totals, key=lambda value: -totals[value][1])[:10]
            rows = [{"label": value, "count": totals[value][0], "weighted_count": round(totals[value][1], 2),
                     "percent": round(100 * totals[value][1] / base, 1) if base else 0} for value in categories]
        if not answered:
            analysis = "Ответов на вопрос пока нет."
        elif numbers:
            analysis = (f"Взвешенное среднее: {numeric_sum / base:.2f}; минимальное значение: {low:g}, "
                        f"максимальное: {high:g}. График показывает распределение по диапазонам.")
        else:
            leaders = sorted((item for item in rows if item["count"]), key=lambda item: -item["weighted_count"])
            leader = leaders[0]
            analysis = (f"Наиболее частый ответ по взвешенной доле: «{leader['label']}» "
                        f"({leader['percent']:.1f}% от ответивших).")
            if question["type"] in ("text", "textarea"):
                analysis += " Показаны до 10 наиболее частых ответов; свободный текст не объединяется по смыслу."
        result["questions"].append({
            "label": question["label"], "type": question["type"], "answered": answered,
            "weighted_base": round(base, 2),
            "mean": round(numeric_sum / base, 2) if base and question["type"] == "number" else None,
            "rows": rows, "analysis": analysis,
        })
    summary = [f"Собрано анкет: {result['count']}; отказов: {result['refusal_count']}. Отказы не включены в расчёт долей."]
    if not result["count"]:
        summary.append("Для выводов пока нет заполненных анкет.")
    else:
        for question in result["questions"][:5]:
            summary.append(f"{question['label']}: {question['analysis']}")
        if len(result["questions"]) > 5:
            summary.append("Остальные вопросы описаны в разделах выше; здесь приведены первые пять.")
    for quota in quotas:
        incomplete = [row["label"] for row in quota["rows"] if row["remaining"]]
        exceeded = [row["label"] for row in quota["rows"] if row["count"] > row["target"]]
        status = (["не выполнены: " + ", ".join(incomplete)] if incomplete else []) + (["превышены: " + ", ".join(exceeded)] if exceeded else [])
        summary.append(f"Квоты «{quota['label']}»: " + ("; ".join(status) if status else "все плановые значения достигнуты") + ".")
    for quota in result["member_quotas"]:
        summary.append(f"Персональная квота «{quota['label']}» ({quota['user_name']}): {quota['count']} из {quota['target']} анкет.")
    if association:
        summary.append(f"Связь «{association['x_label']}» и «{association['y_label']}»: {association['analysis']}")
    summary.append("Выводы описательные: веса и выбранные квоты не позволяют без проверки репрезентативности обобщать результаты на всё население.")
    result["summary"] = summary
    return result


def analytical_snapshot(report):
    """Only publish aggregated, non-small survey cells to the external model."""
    questions = []
    for question in report["questions"]:
        answered = question["answered"]
        item = {"label": question["label"], "type": question["type"],
                "answered": answered if answered >= 5 or not answered else "менее 5"}
        if question["type"] in ("text", "textarea"):
            item["note"] = "Свободные ответы не передаются в ИИ; тематические выводы по этому вопросу недоступны."
            questions.append(item)
            continue
        if question["type"] == "number":
            if answered >= 10:
                item["mean"] = question["mean"]
            else:
                item["note"] = "Числовая база меньше 10: среднее не раскрывается."
            questions.append(item)
            continue
        rows = [row for row in question["rows"] if row["count"]]
        if answered >= 10 and rows and all(row["count"] >= 5 for row in rows):
            item["rows"] = [{"label": row["label"], "count": row["count"], "percent": row["percent"]} for row in rows]
        else:
            item["note"] = ("На вопрос пока нет ответов." if not answered else
                            "Распределение не раскрывается из-за малой базы или категорий с численностью менее пяти.")
        questions.append(item)
    return {"title": report["title"], "goal": report["goal"], "tasks": report["tasks"],
            "count": report["count"], "refusal_count": report["refusal_count"],
            "total_weight": report["total_weight"], "online_count": report["online_count"],
            "public_count": report["public_count"], "excel_count": report["excel_count"],
            "questions": questions, "weighted": report["weighting"] is not None}


def analytical_pdf_bytes(saved):
    candidates = [os.getenv("PDF_TIMES_FONT", ""), os.getenv("PDF_FONT", ""),
                  "C:/Windows/Fonts/times.ttf", "/usr/share/fonts/truetype/msttcorefonts/Times_New_Roman.ttf",
                  "/usr/share/fonts/truetype/dejavu/DejaVuSerif.ttf",
                  "/usr/share/fonts/truetype/liberation2/LiberationSerif-Regular.ttf",
                  "/usr/share/fonts/truetype/noto/NotoSerif-Regular.ttf"]
    font = next((path for path in candidates if path and Path(path).is_file()), None)
    if not font:
        raise ValueError("Для аналитического PDF нужен TTF-шрифт с кириллицей: укажите PDF_TIMES_FONT")
    pdfmetrics.registerFont(TTFont("ReportTimes", font))
    output = io.BytesIO()
    document = SimpleDocTemplate(output, pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm,
                                 topMargin=2 * cm, bottomMargin=2 * cm)
    body = ParagraphStyle("analytical-body", fontName="ReportTimes", fontSize=14, leading=21,
                          firstLineIndent=1.25 * cm, alignment=TA_JUSTIFY, spaceAfter=10)
    heading = ParagraphStyle("analytical-heading", parent=body, firstLineIndent=0, spaceBefore=16, spaceAfter=8)
    snapshot = saved["snapshot"]
    elements = [Paragraph("Аналитический отчёт", heading), Paragraph(escape(snapshot["title"]), heading),
                Paragraph("Дата формирования (UTC): " + escape(saved["generated_at"]) +
                          f". Анкет: {snapshot['count']}; отказов: {snapshot['refusal_count']}; сумма весов: {snapshot['total_weight']:.2f}.", body),
                Paragraph("Подготовлено с помощью ИИ на основе обезличенных агрегированных показателей. "
                          "Не является официальной позицией ведомства. Выводы и рекомендации требуют проверки исследователем; "
                          "без подтверждённой репрезентативности результаты нельзя обобщать на население.", body)]
    if snapshot.get("goal"):
        elements.extend([Paragraph("Цель исследования", heading), Paragraph(escape(snapshot["goal"]), body)])
    if snapshot.get("tasks"):
        elements.extend([Paragraph("Задачи исследования", heading),
                         Paragraph("<br/>".join(escape(line) for line in snapshot["tasks"].splitlines()), body)])
    if saved.get("stale"):
        elements.append(Paragraph("Внимание: данные исследования изменились после создания этой записки. "
                                  "Перед передачей сформируйте новую версию.", body))
    for block in saved["content"].split("\n\n"):
        lines = [line.strip() for line in block.splitlines() if line.strip()]
        if not lines:
            continue
        if len(lines) == 1 and len(lines[0]) < 110 and not lines[0].endswith((".", ";")):
            elements.append(Paragraph(escape(lines[0].lstrip("# ")), heading))
        else:
            elements.append(Paragraph("<br/>".join(escape(line) for line in lines), body))
    elements.append(Paragraph("Статистическая основа отчёта", heading))
    elements.append(Paragraph("Доли взвешены по ответившим на вопрос. Вопросы со свободным текстом и категории "
                              "с численностью менее пяти приведены без раскрытия ответов.", body))
    if not snapshot["questions"]:
        elements.append(Paragraph("Вопросы отсутствовали на момент формирования отчёта.", body))
    for index, question in enumerate(snapshot["questions"], 1):
        elements.append(Paragraph(f"{index}. {escape(question['label'])}", heading))
        elements.append(Paragraph(f"Ответили: {question['answered']}.", body))
        if question["type"] == "number":
            if question.get("mean") is not None:
                elements.append(Paragraph(f"Взвешенное среднее: {question['mean']:.2f}.", body))
        if question.get("note"):
            elements.append(Paragraph(escape(question["note"]), body))
        if question.get("rows"):
            if question["type"] == "multiple":
                elements.append(Paragraph("При множественном выборе сумма долей может превышать 100%.", body))
            values = [[Paragraph("Вариант", heading), Paragraph("Анкет", heading), Paragraph("Доля", heading)]]
            values += [[Paragraph(escape(row["label"]), body), str(row["count"]), f"{row['percent']:.1f}%"] for row in question["rows"]]
            table = Table(values, colWidths=[290, 80, 75], repeatRows=1, hAlign="LEFT")
            table.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"),
                                       ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eaf6ff")),
                                       ("FONTNAME", (0, 0), (-1, -1), "ReportTimes"),
                                       ("FONTSIZE", (0, 0), (-1, -1), 14),
                                       ("BOTTOMPADDING", (0, 0), (-1, -1), 8)]))
            elements.append(table)
    document.build(elements)
    return output.getvalue()


def pdf_bytes(report):
    candidates = [os.getenv("PDF_FONT", ""), "C:/Windows/Fonts/arial.ttf", "C:/Windows/Fonts/segoeui.ttf",
                  "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf"]
    font = next((path for path in candidates if path and Path(path).is_file()), None)
    if not font:
        raise ValueError("Для PDF нужен шрифт с кириллицей: укажите путь в PDF_FONT")
    pdfmetrics.registerFont(TTFont("LabSans", font))
    output = io.BytesIO()
    document = SimpleDocTemplate(output, pagesize=A4, leftMargin=40, rightMargin=40, topMargin=45, bottomMargin=45)
    title = ParagraphStyle("title", fontName="LabSans", fontSize=17, leading=23, textColor=colors.HexColor("#193a54"), spaceAfter=14)
    body = ParagraphStyle("body", fontName="LabSans", fontSize=9, leading=14, textColor=colors.HexColor("#193a54"), spaceAfter=10)
    small = ParagraphStyle("small", parent=body, fontSize=8, leading=11, spaceAfter=0)
    elements = [Paragraph("Отчёт по исследованию: " + escape(report["title"]), title),
                  Paragraph(f"Анкет: {report['count']} (интервью: {report['online_count']}, по ссылке: {report['public_count']}, Excel: {report['excel_count']}). Отказов: {report['refusal_count']}. Сумма весов: {report['total_weight']:.2f}.", body),
                 Paragraph("Доли рассчитаны от суммы весов ответивших на каждый вопрос. Для множественного выбора сумма долей может быть больше 100%.", body)]
    if report["goal"]:
        elements.append(Paragraph("Цель исследования: " + escape(report["goal"]), body))
    if report["tasks"]:
        elements.append(Paragraph("Задачи исследования: " + "<br/>".join(escape(line) for line in report["tasks"].splitlines()), body))
    if report["weighting"]:
        question = next(q for q in report["weighting_questions"] if q["id"] == report["weighting"]["question_id"])
        elements.append(Paragraph("Взвешивание по вопросу «" + escape(question["label"]) + "»: " +
                                  "; ".join(escape(option) + f" = {weight:g}" for option, weight in report["weighting"]["coefficients"].items()) +
                                  ". Для пустых ответов коэффициент равен 1; сохранённые веса анкет умножаются на коэффициенты.", body))
    for index, question in enumerate(report["questions"], 1):
        heading = Paragraph(f"{index}. {escape(question['label'])}", body)
        description = Paragraph(f"Ответили: {question['answered']}; взвешенная база: {question['weighted_base']:.2f}" +
                                (f"; взвешенное среднее: {question['mean']:.2f}" if question["mean"] is not None else ""), small)
        analysis = Paragraph(escape(question["analysis"]), body)
        chart = None
        if question["rows"] and question["answered"]:
            chart_rows = []
            for item in question["rows"]:
                bar = Drawing(225, 12)
                bar.add(Rect(0, 1, 170, 10, fillColor=colors.HexColor("#eaf6ff"), strokeColor=None))
                bar.add(Rect(0, 1, 170 * item["percent"] / 100, 10, fillColor=colors.HexColor("#357db4"), strokeColor=None))
                chart_rows.append([Paragraph(escape(str(item["label"])), small), bar, f"{item['percent']:.1f}%"])
            chart = Table(chart_rows, colWidths=[225, 225, 55], hAlign="LEFT")
            chart.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                                       ("FONTNAME", (0, 0), (-1, -1), "LabSans"), ("FONTSIZE", (0, 0), (-1, -1), 8),
                                       ("BOTTOMPADDING", (0, 0), (-1, -1), 5)]))
        table = None
        if question["rows"]:
            values = [[Paragraph("Вариант", small), Paragraph("Анкет", small), Paragraph("Взвешенно", small), Paragraph("Доля", small)]]
            for item in question["rows"]:
                values.append([Paragraph(escape(str(item["label"])), small), str(item["count"]),
                               f"{item['weighted_count']:.2f}", f"{item['percent']:.1f}%"])
            table = Table(values, colWidths=[260, 55, 90, 55], repeatRows=1, hAlign="LEFT")
            table.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eaf6ff")),
                                       ("LINEBELOW", (0, 0), (-1, -1), .3, colors.HexColor("#dbeaf4")),
                                       ("FONTNAME", (0, 0), (-1, -1), "LabSans"), ("FONTSIZE", (0, 0), (-1, -1), 8),
                                       ("LEFTPADDING", (0, 0), (-1, -1), 6), ("TOPPADDING", (0, 0), (-1, -1), 6),
                                       ("BOTTOMPADDING", (0, 0), (-1, -1), 6)]))
        elements.extend([heading, description, analysis, Spacer(1, 6)])
        if chart:
            elements.extend([chart, Spacer(1, 8)])
        if table:
            elements.append(table)
        elements.append(Spacer(1, 15))
    if report["association"]:
        association = report["association"]
        elements.extend([Paragraph("Связь переменных", title), Paragraph(escape(association["x_label"] + " / " + association["y_label"]), body),
                         Paragraph(escape(association["analysis"]), body)])
        if association["details"] and (association["method"] != "V Крамера" or len(association["columns"]) <= 7):
            if association["method"] == "V Крамера":
                values = [["", *association["columns"]]] + [[row["label"], *[f"{value:.2f}" for value in row["values"]]] for row in association["details"]]
            else:
                values = [["Группа", "Анкет", "Взвешенная база", "Среднее"]] + [[row["label"], str(row["count"]), f"{row['weighted_base']:.2f}", "—" if row["mean"] is None else f"{row['mean']:.2f}"] for row in association["details"]]
            width = min(510 / len(values[0]), 120)
            grid = Table([[Paragraph(escape(str(cell)), small) for cell in row] for row in values], colWidths=[width] * len(values[0]), repeatRows=1, hAlign="LEFT")
            grid.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eaf6ff"))]))
            elements.extend([grid, Spacer(1, 14)])
    elements.append(Paragraph("Итоговые аналитические выводы", title))
    elements.extend(Paragraph(escape(item), body) for item in report["summary"])
    document.build(elements)
    return output.getvalue()
