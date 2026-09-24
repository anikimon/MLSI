import base64
import csv
import io
import json
import os
import sqlite3
import tempfile
import threading
import unittest
from contextlib import closing
from http.cookiejar import CookieJar
from pathlib import Path
from urllib.error import HTTPError
from urllib.request import HTTPCookieProcessor, Request, build_opener
from unittest.mock import patch

from openpyxl import Workbook
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

from presentations import pptx_bytes
from reports import analytical_pdf_bytes
from server import create_server, init_db


class PlannerTest(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.key_file = Path(self.temp.name) / "key.txt"
        key_file_patch = patch("server.DEEPSEEK_KEY_FILE", self.key_file)
        key_file_patch.start()
        self.addCleanup(key_file_patch.stop)
        self.db_path = Path(self.temp.name) / "lab.db"
        self.server = create_server(self.db_path, port=0, seed_demo=False)
        self.thread = threading.Thread(target=self.server.serve_forever, daemon=True)
        self.thread.start()
        self.base = f"http://127.0.0.1:{self.server.server_address[1]}"
        self.admin = build_opener(HTTPCookieProcessor(CookieJar()))
        self.interviewer = build_opener(HTTPCookieProcessor(CookieJar()))
        self.researcher = build_opener(HTTPCookieProcessor(CookieJar()))

    def tearDown(self):
        self.server.shutdown()
        self.server.server_close()
        self.thread.join()
        self.temp.cleanup()

    def request(self, client, path, method="GET", body=None, binary=False):
        payload = json.dumps(body, ensure_ascii=False).encode() if body is not None else None
        request = Request(self.base + "/api" + path, data=payload, method=method,
                          headers={"Content-Type": "application/json"} if body is not None else {})
        try:
            with client.open(request) as response:
                data = response.read()
                return response.status, data if binary else json.loads(data)
        except HTTPError as error:
            return error.code, json.loads(error.read())

    @staticmethod
    def excel(headers, rows):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(headers)
        for row in rows:
            sheet.append(row)
        buffer = io.BytesIO()
        workbook.save(buffer)
        workbook.close()
        return base64.b64encode(buffer.getvalue()).decode("ascii")

    def test_existing_studies_gain_goal_and_tasks(self):
        path = Path(self.temp.name) / "legacy.db"
        with closing(sqlite3.connect(path)) as db, db:
            db.execute("""CREATE TABLE studies (
                id INTEGER PRIMARY KEY, title TEXT NOT NULL, description TEXT NOT NULL DEFAULT '',
                due_date TEXT, stage TEXT NOT NULL DEFAULT 'development',
                responsible_id INTEGER, created_at TEXT NOT NULL)""")
            db.execute("INSERT INTO studies (title, created_at) VALUES ('Старое исследование', '2026-01-01')")
        init_db(path)
        with closing(sqlite3.connect(path)) as db:
            self.assertEqual(db.execute("SELECT goal, tasks FROM studies WHERE id = 1").fetchone(), ("", ""))

    def test_demo_study_is_seeded_once_and_has_385_fictional_responses(self):
        self.server.RequestHandlerClass.demo_enabled = True
        self.assertTrue(self.request(self.admin, "/session")[1]["setup"])
        self.request(self.admin, "/setup", "POST", {"name": "Админ", "email": "admin@test.org", "password": "secure-pass-123"})
        studies = self.request(self.admin, "/studies")[1]["studies"]
        self.assertEqual(len(studies), 1)
        demo = studies[0]
        self.assertIn("ДЕМО · Warhammer 40,000", demo["title"])
        self.assertEqual(demo["response_count"], 385)
        self.assertEqual(len(demo["questions"]), 10)
        sid = demo["id"]
        report = self.request(self.admin, f"/studies/{sid}/report")[1]
        self.assertEqual(report["count"], 385)
        self.assertEqual(report["excel_count"], 385)
        self.assertGreater(len(report["questions"][3]["rows"]), 1)
        self.assertEqual(self.request(self.admin, f"/studies/{sid}", "DELETE")[0], 409)
        with closing(sqlite3.connect(self.db_path)) as db:
            self.assertEqual(db.execute("SELECT COUNT(*) FROM responses WHERE study_id = ?", (sid,)).fetchone()[0], 385)
        restart = create_server(self.db_path, port=0)
        restart.server_close()
        self.assertEqual(len(self.request(self.admin, "/studies")[1]["studies"]), 1)
        self.assertEqual(self.request(self.admin, f"/studies/{sid}/report")[1]["count"], 385)

    def test_demo_is_added_to_existing_database_on_startup(self):
        self.request(self.admin, "/setup", "POST", {"name": "Админ", "email": "admin@test.org", "password": "secure-pass-123"})
        own = self.request(self.admin, "/studies", "POST", {"title": "Моё исследование"})[1]["id"]
        restarted = create_server(self.db_path, port=0)
        restarted.server_close()
        studies = self.request(self.admin, "/studies")[1]["studies"]
        self.assertEqual(len(studies), 2)
        self.assertEqual(next(s for s in studies if s["id"] == own)["title"], "Моё исследование")
        self.assertEqual(next(s for s in studies if s["id"] != own)["response_count"], 385)

    def test_role_boundaries_and_study_deletion(self):
        self.request(self.admin, "/setup", "POST", {"name": "Админ", "email": "admin@test.org", "password": "secure-pass-123"})
        self.request(self.admin, "/users", "POST", {"name": "Исследователь", "email": "research@test.org",
            "password": "secure-pass-456", "role": "researcher"})
        self.request(self.researcher, "/login", "POST", {"email": "research@test.org", "password": "secure-pass-456"})
        created = self.request(self.researcher, "/studies", "POST", {"title": "Удаляемое", "description": "",
            "due_date": "", "responsible_id": None})
        self.assertEqual(created[0], 201)
        sid = created[1]["id"]
        self.assertEqual(self.request(self.researcher, "/users")[0], 403)
        self.assertEqual(self.request(self.researcher, f"/studies/{sid}/editor/guide")[0], 403)
        self.assertEqual(self.request(self.researcher, f"/studies/{sid}/focus", "POST", {"title": "Группа", "guide": "Вопрос"})[0], 201)
        self.assertEqual(self.request(self.researcher, f"/studies/{sid}/responses", "POST", {"answers": {}})[0], 403)
        self.assertEqual(self.request(self.researcher, f"/studies/{sid}", "DELETE")[0], 403)
        self.assertEqual(self.request(self.admin, f"/studies/{sid}/editor/guide", "PUT", {"body": "Черновик"})[0], 200)
        self.assertEqual(self.request(self.admin, f"/studies/{sid}", "DELETE")[0], 200)
        self.assertEqual(self.request(self.admin, f"/studies/{sid}", "DELETE")[0], 404)
        self.assertEqual(self.request(self.admin, "/studies")[1]["studies"], [])
        with closing(sqlite3.connect(self.db_path)) as db:
            for table in ("files", "questionnaires", "responses", "refusals", "weighting", "submissions",
                          "study_messages", "study_quotas", "study_associations", "study_members", "member_quotas",
                           "public_links", "focus_sessions", "editor_drafts", "analytical_reports", "presentations", "presentation_templates"):
                self.assertEqual(db.execute(f"SELECT COUNT(*) FROM {table} WHERE study_id = ?", (sid,)).fetchone()[0], 0)

    def test_calendar_tasks_files_summary_and_permissions(self):
        self.assertEqual(self.request(self.admin, "/tasks?month=2026-09")[0], 401)
        self.request(self.admin, "/setup", "POST", {"name": "Админ", "email": "admin@test.org", "password": "secure-pass-123"})
        self.request(self.admin, "/users", "POST", {"name": "Исследователь", "email": "research@test.org",
            "password": "secure-pass-456", "role": "researcher"})
        self.request(self.admin, "/users", "POST", {"name": "Интервьюер", "email": "field@test.org",
            "password": "secure-pass-789", "role": "interviewer"})
        self.request(self.researcher, "/login", "POST", {"email": "research@test.org", "password": "secure-pass-456"})
        self.request(self.interviewer, "/login", "POST", {"email": "field@test.org", "password": "secure-pass-789"})
        sid = self.request(self.admin, "/studies", "POST", {"title": "Проект", "description": ""})[1]["id"]
        payload = {"title": "Подготовить материалы", "description": "Отправить в отдел", "due_date": "2020-02-29", "study_id": sid}
        self.assertEqual(self.request(self.interviewer, "/tasks?month=2020-02")[0], 403)
        self.assertEqual(self.request(self.interviewer, "/tasks", "POST", payload)[0], 403)
        self.assertEqual(self.request(self.admin, "/tasks", "POST", {**payload, "due_date": "2020-02-30"})[0], 400)
        self.assertEqual(self.request(self.admin, "/tasks", "POST", {**payload, "study_id": 999})[0], 404)
        tid = self.request(self.researcher, "/tasks", "POST", payload)[1]["id"]
        pdf = b"%PDF-1.4\n%%EOF"
        jpeg = b"\xff\xd8\xff\xe0picture\xff\xd9"
        path = f"/tasks/{tid}/files"
        self.assertEqual(self.request(self.researcher, path, "POST", {"name": "wrong.pdf", "content": base64.b64encode(b"abc").decode()})[0], 400)
        self.assertEqual(self.request(self.admin, path, "POST", {"name": "test.png", "content": base64.b64encode(jpeg).decode()})[0], 400)
        self.assertEqual(self.request(self.admin, path, "POST", {"name": "plan.pdf", "content": base64.b64encode(pdf).decode()})[0], 201)
        self.assertEqual(self.request(self.researcher, path, "POST", {"name": "photo.jpg", "content": base64.b64encode(jpeg).decode()})[0], 201)
        status, result = self.request(self.researcher, "/tasks?month=2020-02")
        self.assertEqual(status, 200)
        self.assertEqual(result["summary"], {"total": 1, "completed": 0, "overdue": 1, "files": 2})
        self.assertEqual(result["tasks"][0]["study_title"], "Проект")
        fid = result["tasks"][0]["files"][0]["id"]
        self.assertEqual(self.request(self.interviewer, f"{path}/{fid}")[0], 403)
        self.assertEqual(self.request(self.admin, f"{path}/{fid}", binary=True), (200, pdf))
        self.assertEqual(self.request(self.researcher, f"/tasks/{tid}", "PATCH", {**payload, "status": "done"})[0], 200)
        self.assertEqual(self.request(self.admin, "/tasks?month=2020-02")[1]["summary"]["completed"], 1)
        self.assertEqual(self.request(self.admin, f"/tasks/{tid}", "PATCH", {**payload, "status": "invalid"})[0], 400)
        self.assertEqual(self.request(self.admin, f"{path}/{fid}", "DELETE")[0], 200)
        self.assertEqual(self.request(self.admin, f"/studies/{sid}", "DELETE")[0], 200)
        self.assertIsNone(self.request(self.admin, "/tasks?month=2020-02")[1]["tasks"][0]["study_id"])
        self.assertEqual(self.request(self.admin, f"/tasks/{tid}", "DELETE")[0], 200)
        self.assertEqual(self.request(self.admin, "/tasks?month=2020-02")[1]["summary"]["files"], 0)
        with closing(sqlite3.connect(self.db_path)) as db:
            self.assertEqual(db.execute("SELECT COUNT(*) FROM calendar_files").fetchone()[0], 0)

    def test_project_survey_permissions_and_export(self):
        with self.admin.open(self.base + "/") as response:
            self.assertIn("НИЛ «МЛСИ»", response.read().decode("utf-8"))
        self.assertTrue(self.request(self.admin, "/session")[1]["setup"])
        self.assertEqual(self.request(self.admin, "/setup", "POST", {"name": "Админ", "email": "admin@test.org", "password": "secure-pass-123"})[0], 201)
        self.assertEqual(self.request(self.admin, "/users", "POST", {"name": "Интервьюер", "email": "field@test.org", "password": "secure-pass-456", "role": "interviewer"})[0], 201)
        self.assertEqual(self.request(self.interviewer, "/login", "POST", {"email": "field@test.org", "password": "secure-pass-456"})[0], 200)

        study = {"title": "Городская среда", "description": "Опрос жителей", "goal": "Оценить городскую среду",
                 "tasks": "Измерить оценки\nОписать проблемы", "due_date": "2026-10-10", "responsible_id": 1}
        self.assertEqual(self.request(self.interviewer, "/studies", "POST", study)[0], 403)
        code, result = self.request(self.admin, "/studies", "POST", study)
        self.assertEqual(code, 201)
        study_id = result["id"]
        self.assertEqual(self.request(self.admin, f"/studies/{study_id}", "PATCH", {
            **{key: value for key, value in study.items() if key not in ("goal", "tasks")}, "stage": "field"})[0], 200)

        attachment = base64.b64encode(b"brief contents").decode()
        self.assertEqual(self.request(self.admin, f"/studies/{study_id}/files", "POST", {"name": "brief.txt", "kind": "brief", "content": attachment})[0], 201)
        project = self.request(self.interviewer, "/studies")[1]["studies"][0]
        self.assertEqual(set(project), {"id", "title", "stage", "questions"})
        self.assertEqual(project["stage"], "field")
        self.assertEqual(self.request(self.interviewer, "/users")[0], 403)
        admin_project = self.request(self.admin, "/studies")[1]["studies"][0]
        self.assertEqual(admin_project["goal"], study["goal"])
        self.assertEqual(admin_project["tasks"], study["tasks"])
        file_id = admin_project["files"][0]["id"]
        self.assertEqual(self.request(self.interviewer, f"/studies/{study_id}/files/{file_id}", binary=True)[0], 403)
        self.assertEqual(self.request(self.admin, f"/studies/{study_id}/files/{file_id}", binary=True)[1], b"brief contents")

        questions = [{"label": "Оценка района", "type": "single", "options": ["Хорошо", "Плохо"], "required": True},
                     {"label": "Комментарий", "type": "text", "required": False}]
        self.assertEqual(self.request(self.admin, f"/studies/{study_id}/questionnaire", "POST", {"questions": questions})[0], 200)
        self.assertEqual(self.request(self.interviewer, f"/studies/{study_id}/responses", "POST", {"code": "А1", "answers": {"q1": "Не существует", "q2": ""}})[0], 400)
        self.assertEqual(self.request(self.interviewer, f"/studies/{study_id}/responses", "POST", {"code": "=1+1", "weight": 2, "answers": {"q1": "Хорошо", "q2": "+опасное значение"}})[0], 201)
        self.assertEqual(self.request(self.interviewer, f"/studies/{study_id}/responses", "POST", {"code": "x", "weight": 0, "answers": {"q1": "Хорошо"}})[0], 400)
        self.assertEqual(self.request(self.interviewer, f"/studies/{study_id}/responses.csv")[0], 403)
        self.assertEqual(self.request(self.interviewer, f"/studies/{study_id}/report")[0], 403)
        self.assertEqual(self.request(self.admin, f"/studies/{study_id}/questionnaire", "POST", {"questions": questions})[0], 409)
        bad = self.excel(["code", "weight", "q1", "q2"], [["r1", 3, "Хорошо", ""], ["r2", 1, "Нет варианта", ""]])
        self.assertEqual(self.request(self.admin, f"/studies/{study_id}/import", "POST", {"name": "bad.xlsx", "content": bad})[0], 400)
        self.assertEqual(self.request(self.admin, "/studies")[1]["studies"][0]["response_count"], 1)
        content = self.excel(["code", "weight", "q1", "q2"], [["r1", 3, "Хорошо", ""], ["r2", 1, "Плохо", "комментарий"]])
        imported = self.request(self.admin, f"/studies/{study_id}/import", "POST", {"name": "responses.xlsx", "content": content})
        self.assertEqual(imported[1]["imported"], 2)
        report = self.request(self.admin, f"/studies/{study_id}/report")[1]
        self.assertEqual((report["count"], report["online_count"], report["excel_count"], report["total_weight"]), (3, 1, 2, 6))
        self.assertEqual(report["questions"][0]["rows"][0]["weighted_count"], 5)
        self.assertEqual(report["questions"][0]["rows"][0]["percent"], 83.3)
        self.assertIn("Хорошо", report["questions"][0]["analysis"])
        self.assertIn("свободный текст", report["questions"][1]["analysis"])
        pdf_status, pdf = self.request(self.admin, f"/studies/{study_id}/report.pdf", binary=True)
        self.assertEqual(pdf_status, 200)
        self.assertTrue(pdf.startswith(b"%PDF-"))
        status, content = self.request(self.admin, f"/studies/{study_id}/responses.csv", binary=True)
        self.assertEqual(status, 200)
        rows = list(csv.reader(io.StringIO(content.decode("utf-8-sig"))))
        self.assertEqual(rows[1][2], "'=1+1")
        self.assertEqual(rows[1][5], "2.0")
        self.assertEqual(rows[1][-1], "'+опасное значение")
        self.assertEqual(self.request(self.admin, "/studies")[1]["studies"][0]["response_count"], 3)
        self.assertEqual(self.request(self.interviewer, "/logout", "POST", {})[0], 200)
        self.assertEqual(self.request(self.interviewer, "/studies")[0], 401)

    def test_excel_without_questionnaire_creates_report(self):
        self.request(self.admin, "/setup", "POST", {"name": "Админ", "email": "admin@test.org", "password": "secure-pass-123"})
        _, created = self.request(self.admin, "/studies", "POST", {"title": "Экспресс-опрос", "description": "", "due_date": "", "responsible_id": None})
        study_id = created["id"]
        content = self.excel(["code", "weight", "Возраст", "Район"], [["a", 2, 20, "Север"], ["b", 1, 50, "Юг"]])
        status, result = self.request(self.admin, f"/studies/{study_id}/import", "POST", {"name": "data.xlsx", "content": content})
        self.assertEqual(status, 201)
        self.assertTrue(result["created_questionnaire"])
        report = self.request(self.admin, f"/studies/{study_id}/report")[1]
        self.assertEqual(report["questions"][0]["type"], "number")
        self.assertEqual(report["questions"][0]["mean"], 30)
        self.assertEqual(sum(row["count"] for row in report["questions"][0]["rows"]), 2)
        self.assertEqual([row["percent"] for row in report["questions"][0]["rows"] if row["count"]], [66.7, 33.3])
        self.assertIn("минимальное значение: 20", report["questions"][0]["analysis"])
        self.assertEqual(report["questions"][1]["rows"][0]["percent"], 66.7)
        self.assertTrue(self.request(self.admin, f"/studies/{study_id}/report.pdf", binary=True)[1].startswith(b"%PDF-"))

    def test_existing_responses_get_default_weight_and_source(self):
        old = Path(self.temp.name) / "old.db"
        with closing(sqlite3.connect(old)) as db:
            with db:
                db.execute("""CREATE TABLE responses (
                    id INTEGER PRIMARY KEY, study_id INTEGER, interviewer_id INTEGER,
                    code TEXT NOT NULL, answers TEXT NOT NULL, created_at TEXT NOT NULL)""")
                db.execute("INSERT INTO responses VALUES (1, 1, 1, 'old', '{}', '2026-01-01')")
        init_db(old)
        with closing(sqlite3.connect(old)) as db:
            self.assertEqual(db.execute("SELECT weight, source FROM responses WHERE id = 1").fetchone(), (1, "online"))

    def test_multiple_choices_and_missing_answers_use_answered_base(self):
        self.request(self.admin, "/setup", "POST", {"name": "Админ", "email": "admin@test.org", "password": "secure-pass-123"})
        _, created = self.request(self.admin, "/studies", "POST", {"title": "Проверка весов", "description": "", "due_date": "", "responsible_id": None})
        study_id = created["id"]
        questions = [{"label": "Выбор", "type": "multiple", "options": ["A", "B"], "required": False},
                     {"label": "Баллы", "type": "number", "required": False}]
        self.request(self.admin, f"/studies/{study_id}/questionnaire", "POST", {"questions": questions})
        for weight, choices, score in [(2, ["A", "B"], "10"), (1, ["A"], "20"), (3, [], "")]:
            status, _ = self.request(self.admin, f"/studies/{study_id}/responses", "POST",
                                     {"code": "", "weight": weight, "answers": {"q1": choices, "q2": score}})
            self.assertEqual(status, 201)
        report = self.request(self.admin, f"/studies/{study_id}/report")[1]
        self.assertEqual(report["total_weight"], 6)
        self.assertEqual(report["questions"][0]["weighted_base"], 3)
        self.assertEqual([row["percent"] for row in report["questions"][0]["rows"]], [100, 66.7])
        self.assertEqual(report["questions"][1]["mean"], 13.33)
        self.assertEqual(sum(row["count"] for row in report["questions"][1]["rows"]), 2)
        self.assertEqual(sum(row["weighted_count"] for row in report["questions"][1]["rows"]), 3)
        self.assertIn("Взвешенное среднее: 13.33", report["questions"][1]["analysis"])

    def test_report_without_answers_and_constant_number(self):
        self.request(self.admin, "/setup", "POST", {"name": "Админ", "email": "admin@test.org", "password": "secure-pass-123"})
        _, created = self.request(self.admin, "/studies", "POST", {"title": "Оценки", "description": "", "due_date": "", "responsible_id": None})
        study_id = created["id"]
        self.request(self.admin, f"/studies/{study_id}/questionnaire", "POST", {"questions": [
            {"label": "Баллы", "type": "number", "required": False},
            {"label": "Замечания", "type": "text", "required": False}]})
        report = self.request(self.admin, f"/studies/{study_id}/report")[1]
        self.assertEqual(report["questions"][0]["analysis"], "Ответов на вопрос пока нет.")
        self.assertEqual(report["questions"][0]["rows"], [])
        self.request(self.admin, f"/studies/{study_id}/responses", "POST", {"answers": {"q1": "5", "q2": ""}})
        report = self.request(self.admin, f"/studies/{study_id}/report")[1]
        self.assertEqual(report["questions"][0]["rows"][0]["percent"], 100)
        self.assertEqual(report["questions"][0]["rows"][0]["label"], "5")
        self.assertEqual(report["questions"][1]["analysis"], "Ответов на вопрос пока нет.")
        self.assertTrue(self.request(self.admin, f"/studies/{study_id}/report.pdf", binary=True)[1].startswith(b"%PDF-"))

    def test_group_weighting_refusals_and_offline_retries(self):
        self.request(self.admin, "/setup", "POST", {"name": "Админ", "email": "admin@test.org", "password": "secure-pass-123"})
        self.request(self.admin, "/users", "POST", {"name": "Полевой", "email": "field@test.org", "password": "secure-pass-456", "role": "interviewer"})
        self.request(self.interviewer, "/login", "POST", {"email": "field@test.org", "password": "secure-pass-456"})
        _, created = self.request(self.admin, "/studies", "POST", {"title": "Поле", "description": "", "due_date": "", "responsible_id": None})
        study_id = created["id"]
        self.request(self.admin, f"/studies/{study_id}/questionnaire", "POST", {"questions": [
            {"label": "Группа", "type": "single", "options": ["A", "B"], "required": False},
            {"label": "Оценка", "type": "number", "required": True}]})
        rule_url = f"/studies/{study_id}/weighting"
        self.assertEqual(self.request(self.interviewer, rule_url, "PUT", {"question_id": "q1", "coefficients": {"A": 2, "B": 1}})[0], 403)
        for rule in ({"question_id": "q2", "coefficients": {}},
                     {"question_id": "q1", "coefficients": {"A": 2}},
                     {"question_id": "q1", "coefficients": {"A": 0, "B": 1}}):
            self.assertEqual(self.request(self.admin, rule_url, "PUT", rule)[0], 400)
        self.assertEqual(self.request(self.admin, rule_url, "PUT", {"question_id": "q1", "coefficients": {"A": 2, "B": 0.5}})[0], 200)
        answer = {"client_id": "12345678-1234-1234-1234-123456789abc", "answers": {"q1": "A", "q2": "10"}}
        self.assertEqual(self.request(self.interviewer, f"/studies/{study_id}/responses", "POST", answer)[0], 201)
        self.assertEqual(self.request(self.interviewer, f"/studies/{study_id}/responses", "POST", answer)[0], 200)
        self.assertEqual(self.request(self.admin, f"/studies/{study_id}/responses", "POST", answer)[0], 409)
        self.assertEqual(self.request(self.interviewer, f"/studies/{study_id}/responses", "POST", {**answer, "answers": {"q1": "B", "q2": "10"}})[0], 409)
        self.assertEqual(self.request(self.admin, f"/studies/{study_id}/questionnaire", "POST", {"questions": []})[0], 409)
        self.request(self.admin, f"/studies/{study_id}/responses", "POST", {"answers": {"q1": "B", "q2": "20"}, "weight": 2})
        self.request(self.admin, f"/studies/{study_id}/responses", "POST", {"answers": {"q1": "", "q2": "30"}})
        report = self.request(self.admin, f"/studies/{study_id}/report")[1]
        self.assertEqual(report["total_weight"], 4)
        self.assertEqual([row["weighted_count"] for row in report["questions"][0]["rows"]], [2, 1])
        self.assertEqual(report["questions"][1]["mean"], 17.5)
        self.assertEqual(report["weighting"]["question_id"], "q1")
        refusal = {"client_id": "12345678-1234-1234-1234-123456789abd"}
        self.assertEqual(self.request(self.interviewer, f"/studies/{study_id}/refusals", "POST", refusal)[0], 201)
        self.assertEqual(self.request(self.interviewer, f"/studies/{study_id}/refusals", "POST", refusal)[0], 200)
        self.assertEqual(self.request(self.admin, "/studies")[1]["studies"][0]["refusal_count"], 1)
        self.assertEqual(self.request(self.admin, f"/studies/{study_id}/report")[1]["refusal_count"], 1)
        self.assertTrue(self.request(self.admin, f"/studies/{study_id}/report.pdf", binary=True)[1].startswith(b"%PDF-"))
        self.assertEqual(self.request(self.admin, rule_url, "PUT", {"question_id": None})[0], 200)
        self.assertEqual(self.request(self.admin, f"/studies/{study_id}/report")[1]["total_weight"], 4)

    def test_mobile_assets_and_offline_queue_support(self):
        for path, content_type in (("/manifest.webmanifest", "application/manifest+json"),
                                   ("/sw.js", "text/javascript"), ("/icon.svg", "image/svg+xml")):
            with self.admin.open(self.base + path) as response:
                self.assertEqual(response.status, 200)
                self.assertIn(content_type, response.headers["Content-Type"])
        with self.admin.open(self.base + "/") as response:
            page = response.read().decode("utf-8")
            self.assertIn("/manifest.webmanifest", page)
            self.assertIn("record-refusal", page)
            self.assertNotIn('name="weight"', page)

    def test_study_chat_quotas_and_associations(self):
        self.request(self.admin, "/setup", "POST", {"name": "Админ", "email": "admin@test.org", "password": "secure-pass-123"})
        self.request(self.admin, "/users", "POST", {"name": "Полевой", "email": "field@test.org", "password": "secure-pass-456", "role": "interviewer"})
        self.request(self.interviewer, "/login", "POST", {"email": "field@test.org", "password": "secure-pass-456"})
        _, created = self.request(self.admin, "/studies", "POST", {"title": "Квоты", "description": "", "due_date": "", "responsible_id": None})
        study_id = created["id"]
        self.request(self.admin, f"/studies/{study_id}/questionnaire", "POST", {"questions": [
            {"label": "Пол", "type": "single", "options": ["Ж", "М"], "required": True},
            {"label": "Возраст", "type": "number", "required": True},
            {"label": "Район", "type": "single", "options": ["Центр", "Север"], "required": True},
            {"label": "Способ", "type": "multiple", "options": ["Авто", "Пешком"], "required": False}]})
        quotas = f"/studies/{study_id}/quotas"
        self.assertEqual(self.request(self.interviewer, quotas, "PUT", {"question_id": "q1", "rows": []})[0], 403)
        for invalid in ([{"option": "Неизвестно", "target": 1}], [{"option": "Ж", "target": -1}],
                        [{"option": "Ж", "target": 1}, {"option": "Ж", "target": 2}]):
            self.assertEqual(self.request(self.admin, quotas, "PUT", {"question_id": "q1", "rows": invalid})[0], 400)
        self.assertEqual(self.request(self.admin, quotas, "PUT", {"question_id": "q2", "rows": [
            {"min": 18, "max": 25, "target": 2}, {"min": 24, "max": 35, "target": 1} ]})[0], 400)
        self.assertEqual(self.request(self.admin, quotas, "PUT", {"question_id": "q1", "rows": [
            {"option": "Ж", "target": 2}, {"option": "М", "target": 1}]})[0], 200)
        self.assertEqual(self.request(self.admin, quotas, "PUT", {"question_id": "q2", "rows": [
            {"min": 18, "max": 25, "target": 1}, {"min": 25, "max": 35, "target": 2}]})[0], 200)
        self.assertEqual(self.request(self.admin, quotas, "PUT", {"question_id": "q3", "rows": [
            {"option": "Центр", "target": 1}]})[0], 200)
        self.assertEqual(self.request(self.admin, quotas, "PUT", {"question_id": "q4", "rows": [
            {"option": "Авто", "target": 2}]})[0], 200)
        self.assertEqual(self.request(self.admin, f"/studies/{study_id}/questionnaire", "POST", {"questions": []})[0], 409)
        for sex, age, district, mode in [("Ж", "24", "Центр", ["Авто", "Пешком"]),
                                         ("Ж", "25", "Центр", ["Авто"]), ("М", "30", "Север", [])]:
            self.assertEqual(self.request(self.interviewer, f"/studies/{study_id}/responses", "POST", {"answers": {
                "q1": sex, "q2": age, "q3": district, "q4": mode}})[0], 201)
        listed = self.request(self.admin, "/studies")[1]["studies"][0]["quotas"]
        self.assertEqual([row["count"] for row in listed[0]["rows"]], [2, 1])
        self.assertEqual([row["remaining"] for row in listed[1]["rows"]], [0, 0])
        self.assertEqual(listed[2]["rows"][0]["count"], 2)
        self.assertEqual(listed[3]["rows"][0]["count"], 2)
        other = self.request(self.admin, "/studies", "POST", {"title": "Второе", "description": "", "due_date": "", "responsible_id": None})[1]["id"]
        self.assertEqual(self.request(self.admin, f"/studies/{other}/messages")[1]["messages"], [])
        self.assertEqual(self.request(self.interviewer, f"/studies/{study_id}/messages", "POST", {"body": "  <Полевой> готов  "})[0], 403)
        self.assertEqual(self.request(self.admin, f"/studies/{study_id}/messages", "POST", {"body": "  <Админ> готов  "})[0], 201)
        self.assertEqual(self.request(self.admin, f"/studies/{study_id}/messages", "POST", {"body": "  "})[0], 400)
        messages = self.request(self.admin, f"/studies/{study_id}/messages")[1]["messages"]
        self.assertEqual((messages[0]["author"], messages[0]["body"]), ("Админ", "<Админ> готов"))
        path = f"/studies/{study_id}/associations"
        self.assertEqual(self.request(self.interviewer, path, "PUT", {"x_id": "q1", "y_id": "q3"})[0], 403)
        self.assertEqual(self.request(self.admin, path, "PUT", {"x_id": "q1", "y_id": "q4"})[0], 400)
        self.assertEqual(self.request(self.admin, path, "PUT", {"x_id": "q1", "y_id": "q3"})[0], 200)
        report = self.request(self.admin, f"/studies/{study_id}/report")[1]
        self.assertEqual((report["association"]["method"], report["association"]["value"]), ("V Крамера", 1))
        self.assertTrue(any("Связь" in line for line in report["summary"]))
        self.assertEqual(self.request(self.admin, path, "PUT", {"x_id": "q2", "y_id": "q1"})[0], 200)
        self.assertEqual(self.request(self.admin, f"/studies/{study_id}/report")[1]["association"]["method"], "Отношение корреляции η")
        self.assertEqual(self.request(self.admin, f"/studies/{study_id}/report.pdf", binary=True)[0], 200)
        self.assertEqual(self.request(self.admin, path, "PUT", {"x_id": None, "y_id": None})[0], 200)
        self.assertIsNone(self.request(self.admin, f"/studies/{study_id}/report")[1]["association"])
        self.assertEqual(self.request(self.admin, quotas, "PUT", {"question_id": "q3", "rows": []})[0], 200)

    def test_weighted_pearson_and_constant_pair(self):
        self.request(self.admin, "/setup", "POST", {"name": "Админ", "email": "admin@test.org", "password": "secure-pass-123"})
        study_id = self.request(self.admin, "/studies", "POST", {"title": "Связь", "description": "", "due_date": "", "responsible_id": None})[1]["id"]
        self.request(self.admin, f"/studies/{study_id}/questionnaire", "POST", {"questions": [
            {"label": "X", "type": "number"}, {"label": "Y", "type": "number"}, {"label": "Одинаковые", "type": "number"}]})
        self.request(self.admin, f"/studies/{study_id}/associations", "PUT", {"x_id": "q1", "y_id": "q2"})
        empty = self.request(self.admin, f"/studies/{study_id}/report")[1]
        self.assertIsNone(empty["association"]["value"])
        self.assertTrue(any("нет заполненных анкет" in line for line in empty["summary"]))
        for x, y, weight in [("1", "3", 1), ("2", "2", 2), ("3", "1", 1)]:
            self.request(self.admin, f"/studies/{study_id}/responses", "POST", {"answers": {"q1": x, "q2": y, "q3": "5"}, "weight": weight})
        report = self.request(self.admin, f"/studies/{study_id}/report")[1]
        self.assertEqual((report["association"]["method"], report["association"]["value"]), ("Пирсон r", -1))
        self.assertEqual(report["association"]["weighted_base"], 4)
        self.request(self.admin, f"/studies/{study_id}/associations", "PUT", {"x_id": "q1", "y_id": "q3"})
        self.assertIsNone(self.request(self.admin, f"/studies/{study_id}/report")[1]["association"]["value"])

    def test_members_compound_quotas_and_public_survey(self):
        self.request(self.admin, "/setup", "POST", {"name": "Админ", "email": "admin@test.org", "password": "secure-pass-123"})
        self.request(self.admin, "/users", "POST", {"name": "Оля", "email": "olya@test.org", "password": "secure-pass-456", "role": "interviewer"})
        self.request(self.interviewer, "/login", "POST", {"email": "olya@test.org", "password": "secure-pass-456"})
        sid = self.request(self.admin, "/studies", "POST", {"title": "Город", "description": "", "due_date": "", "responsible_id": None})[1]["id"]
        base = f"/studies/{sid}"
        self.request(self.admin, base + "/questionnaire", "POST", {"questions": [
            {"label": "Пол", "type": "single", "options": ["Ж", "М"]}, {"label": "Возраст", "type": "number"},
            {"label": "Район", "type": "multiple", "options": ["Центр", "Юг"]}]})
        self.assertEqual(self.request(self.interviewer, base + "/members", "PUT", {"user_ids": [2]})[0], 403)
        self.assertEqual(self.request(self.admin, base + "/members", "PUT", {"user_ids": [2, 999]})[0], 400)
        self.assertEqual(self.request(self.admin, base + "/members", "PUT", {"user_ids": [2]})[0], 200)
        path = base + "/member-quotas"
        quota = {"user_id": 2, "label": "Женщины 35+ из центра", "target": 2, "conditions": [
            {"question_id": "q1", "option": "Ж"}, {"question_id": "q2", "min": 35},
            {"question_id": "q3", "option": "Центр"}]}
        self.assertEqual(self.request(self.admin, path, "POST", {**quota, "conditions": [{"question_id": "q2", "min": None, "max": None}]})[0], 400)
        self.assertEqual(self.request(self.admin, path, "POST", {**quota, "conditions": [
            {"question_id": "q1", "option": "Ж"}, {"question_id": "q1", "option": "М"}]})[0], 400)
        self.assertEqual(self.request(self.interviewer, path, "POST", quota)[0], 403)
        self.assertEqual(self.request(self.admin, path, "POST", quota)[0], 201)
        self.assertEqual(self.request(self.admin, base + "/members", "PUT", {"user_ids": []})[0], 409)
        for sex, age, district in [("Ж", "35", ["Центр"]), ("Ж", "34", ["Центр"]), ("М", "40", ["Центр"]),
                                    ("Ж", "50", ["Юг"]), ("Ж", "40", ["Юг", "Центр"])]:
            self.assertEqual(self.request(self.interviewer, base + "/responses", "POST", {"answers": {
                "q1": sex, "q2": age, "q3": district}})[0], 201)
        listed = self.request(self.admin, "/studies")[1]["studies"][0]
        self.assertEqual(listed["member_ids"], [2])
        self.assertEqual((listed["member_quotas"][0]["count"], listed["member_quotas"][0]["remaining"]), (2, 0))
        self.assertEqual(self.request(self.interviewer, base + "/survey-link", "PUT")[0], 403)
        token = self.request(self.admin, base + "/survey-link", "PUT")[1]["token"]
        self.assertEqual(self.request(self.admin, base + "/survey-link")[1]["token"], token)
        self.assertEqual(self.request(self.admin, base + "/questionnaire", "POST", {"questions": []})[0], 409)
        public = build_opener(HTTPCookieProcessor(CookieJar()))
        with public.open(self.base + f"/s/{token}") as page:
            self.assertIn("Онлайн-опрос", page.read().decode("utf-8"))
        self.assertEqual(self.request(public, f"/public/{token}")[1]["title"], "Город")
        self.assertEqual(self.request(public, base + "/report")[0], 401)
        answer = {"client_id": "12345678-1234-1234-1234-123456789abc", "answers": {"q1": "Ж", "q2": "36", "q3": ["Центр"]}}
        self.assertEqual(self.request(public, f"/public/{token}", "POST", answer)[0], 201)
        self.assertEqual(self.request(public, f"/public/{token}", "POST", answer)[0], 200)
        self.assertEqual(self.request(public, f"/public/{token}", "POST", {**answer, "answers": {"q1": "М", "q2": "36", "q3": []}})[0], 409)
        report = self.request(self.admin, base + "/report")[1]
        self.assertEqual((report["count"], report["public_count"], report["member_quotas"][0]["count"]), (6, 1, 2))
        self.assertTrue(self.request(self.admin, base + "/report.pdf", binary=True)[1].startswith(b"%PDF-"))
        self.assertIn("По ссылке", self.request(self.admin, base + "/responses.csv", binary=True)[1].decode("utf-8-sig"))
        self.assertEqual(self.request(self.admin, base + "/survey-link", "DELETE")[0], 200)
        self.assertEqual(self.request(public, f"/public/{token}")[0], 404)
        quota_id = listed["member_quotas"][0]["id"]
        self.assertEqual(self.request(self.admin, path + f"/{quota_id}", "DELETE")[0], 200)
        self.assertEqual(self.request(self.admin, base + "/members", "PUT", {"user_ids": []})[0], 200)

    def test_focus_group_guide_audio_and_transcript(self):
        self.request(self.admin, "/setup", "POST", {"name": "Админ", "email": "admin@test.org", "password": "secure-pass-123"})
        self.request(self.admin, "/users", "POST", {"name": "Модератор", "email": "moderator@test.org", "password": "secure-pass-456", "role": "interviewer"})
        self.request(self.interviewer, "/login", "POST", {"email": "moderator@test.org", "password": "secure-pass-456"})
        sid = self.request(self.admin, "/studies", "POST", {"title": "Фокус-группы", "description": "", "due_date": "", "responsible_id": None})[1]["id"]
        base = f"/studies/{sid}/focus"
        self.assertEqual(self.request(self.interviewer, base, "POST", {"title": "Чужая", "guide": "Текст"})[0], 403)
        created = self.request(self.admin, base, "POST", {"title": "Группа №1", "guide": "Вступление\nВопросы\nЗаключение"})
        self.assertEqual(created[0], 201)
        fid = created[1]["id"]
        self.assertEqual(self.request(self.admin, base)[1]["sessions"][0]["guide"], "Вступление\nВопросы\nЗаключение")
        audio = base64.b64encode(b"\x1a\x45\xdf\xa3example").decode()
        self.assertEqual(self.request(self.interviewer, f"{base}/{fid}/audio", "PUT", {"mime": "audio/webm", "content": audio})[0], 403)
        self.assertEqual(self.request(self.admin, f"/studies/{sid}/members", "PUT", {"user_ids": [2]})[0], 200)
        self.assertEqual(self.request(self.admin, f"{base}/{fid}/audio", "PUT", {"mime": "text/plain", "content": audio})[0], 400)
        self.assertEqual(self.request(self.interviewer, f"{base}/{fid}/audio", "PUT", {"mime": "audio/webm", "content": audio})[0], 403)
        self.assertEqual(self.request(self.interviewer, base)[0], 403)
        self.assertEqual(self.request(self.admin, f"{base}/{fid}/audio", "PUT", {"mime": "audio/webm", "content": audio})[0], 200)
        self.assertEqual(self.request(self.admin, f"{base}/{fid}/audio", binary=True)[1], b"\x1a\x45\xdf\xa3example")
        with patch("server.transcribe_audio", return_value="Расшифрованный разговор") as transcribe:
            self.assertEqual(self.request(self.admin, f"{base}/{fid}/transcribe", "POST", {})[1]["transcript"], "Расшифрованный разговор")
            transcribe.assert_called_once()
        with patch("server.transcribe_audio", return_value=""):
            self.assertEqual(self.request(self.admin, f"{base}/{fid}/transcribe", "POST", {})[0], 422)
        self.assertEqual(self.request(self.admin, base)[1]["sessions"][0]["transcript"], "Расшифрованный разговор")
        self.assertEqual(self.request(self.interviewer, f"{base}/{fid}/transcript", "PUT", {"transcript": "Исправленный текст"})[0], 403)
        self.assertEqual(self.request(self.admin, f"{base}/{fid}/transcript", "PUT", {"transcript": "Исправленный текст"})[0], 200)
        self.assertEqual(self.request(self.admin, base)[1]["sessions"][0]["transcript"], "Исправленный текст")

    def test_editor_drafts_guides_and_ai_proxy(self):
        self.request(self.admin, "/setup", "POST", {"name": "Админ", "email": "admin@test.org", "password": "secure-pass-123"})
        self.request(self.admin, "/users", "POST", {"name": "Модератор", "email": "moderator@test.org", "password": "secure-pass-456", "role": "interviewer"})
        self.request(self.interviewer, "/login", "POST", {"email": "moderator@test.org", "password": "secure-pass-456"})
        sid = self.request(self.admin, "/studies", "POST", {"title": "Исследование", "description": "", "due_date": "", "responsible_id": None})[1]["id"]
        base = f"/studies/{sid}/editor"
        self.assertEqual(self.request(self.admin, base + "/questionnaire")[1]["body"], "")
        self.assertEqual(self.request(self.interviewer, base + "/questionnaire", "PUT", {"body": "Текст"})[0], 403)
        self.assertEqual(self.request(self.admin, base + "/questionnaire", "PUT", {"body": "Как вы добираетесь?\n- Пешком\n- Автобус"})[0], 200)
        self.assertIn("Автобус", self.request(self.admin, base + "/questionnaire")[1]["body"])
        self.assertEqual(self.request(self.admin, base + "/guide", "PUT", {"body": "Вступление\nОбсуждение"})[0], 200)
        self.assertEqual(self.request(self.admin, base + "/guide")[1]["body"], "Вступление\nОбсуждение")
        focus = f"/studies/{sid}/focus"
        fid = self.request(self.admin, focus, "POST", {"title": "Встреча", "guide": "Старый гайд"})[1]["id"]
        self.assertEqual(self.request(self.interviewer, f"{focus}/{fid}/guide", "PUT", {"guide": "Новый гайд"})[0], 403)
        self.assertEqual(self.request(self.admin, f"{focus}/{fid}/guide", "PUT", {"guide": ""})[0], 400)
        self.assertEqual(self.request(self.admin, f"{focus}/{fid}/guide", "PUT", {"guide": "Новый гайд"})[0], 200)
        self.assertEqual(self.request(self.admin, focus)[1]["sessions"][0]["guide"], "Новый гайд")
        question = {"kind": "questionnaire", "body": "Как вы добираетесь?", "question": "Как улучшить вопрос?", "history": []}
        self.assertEqual(self.request(self.interviewer, base + "/assist", "POST", question)[0], 403)
        with patch.dict(os.environ, {"DEEPSEEK_API_KEY": ""}):
            self.assertEqual(self.request(self.admin, base + "/assist", "POST", question)[0], 503)
        self.key_file.write_text("\ufeff file-placeholder \n", encoding="utf-8")
        with patch.dict(os.environ, {"DEEPSEEK_API_KEY": "wrong-env-key"}), patch("server.urlopen") as urlopen:
            urlopen.return_value.__enter__.return_value.read.return_value = b'{"choices":[{"message":{"content":"OK"}}]}'
            self.assertEqual(self.request(self.admin, base + "/assist", "POST", question)[0], 200)
            self.assertEqual(urlopen.call_args.args[0].get_header("Authorization"), "Bearer file-placeholder")
        self.key_file.unlink()
        with patch.dict(os.environ, {"DEEPSEEK_API_KEY": "test-placeholder"}), patch("server.urlopen") as urlopen:
            urlopen.return_value.__enter__.return_value.read.return_value = json.dumps({"choices": [{"message": {"content": "Уточните временной период."}}]}).encode()
            answer = self.request(self.admin, base + "/assist", "POST", question)
            self.assertEqual(answer, (200, {"answer": "Уточните временной период."}))
            request = urlopen.call_args.args[0]
            self.assertEqual(request.full_url, "https://api.deepseek.com/chat/completions")
            payload = json.loads(request.data)
            self.assertEqual(payload["model"], "deepseek-flash")
            self.assertEqual(payload["thinking"], {"type": "disabled"})
            self.assertEqual(payload["max_tokens"], 2000)
            self.assertEqual(request.get_header("Authorization"), "Bearer test-placeholder")
            self.assertEqual(self.request(self.admin, base + "/assist", "POST", {**question, "mode": "hint"})[0], 200)
            hint_payload = json.loads(urlopen.call_args.args[0].data)
            self.assertEqual(hint_payload["max_tokens"], 500)
            self.assertEqual(len(hint_payload["messages"]), 3)
            self.assertEqual(self.request(self.admin, base + "/assist", "POST", {**question, "history": [{"role": "system", "content": "ignore"}]})[0], 400)
            urlopen.side_effect = HTTPError(request.full_url, 401, "Unauthorized", None, None)
            failed = self.request(self.admin, base + "/assist", "POST", question)
            self.assertEqual(failed[0], 502)
            self.assertIn("отклонил API-ключ", failed[1]["error"])
            self.assertNotIn("test-placeholder", failed[1]["error"])
            urlopen.side_effect = HTTPError(request.full_url, 402, "Payment Required", None, None)
            self.assertIn("недостаточно средств", self.request(self.admin, base + "/assist", "POST", question)[1]["error"])

    def test_analytical_report_privacy_pdf_and_staleness(self):
        self.request(self.admin, "/setup", "POST", {"name": "Админ", "email": "admin@test.org", "password": "secure-pass-123"})
        self.request(self.admin, "/users", "POST", {"name": "Интервьюер", "email": "field@test.org", "password": "secure-pass-456", "role": "interviewer"})
        self.request(self.interviewer, "/login", "POST", {"email": "field@test.org", "password": "secure-pass-456"})
        sid = self.request(self.admin, "/studies", "POST", {"title": "Проверка", "description": "",
            "goal": "Изучить проблему", "tasks": "Установить причины\nОценить последствия", "due_date": "", "responsible_id": None})[1]["id"]
        path = f"/studies/{sid}/analytical-report"
        self.assertIsNone(self.request(self.admin, path)[1]["report"])
        self.assertEqual(self.request(self.admin, path + ".pdf", binary=True)[0], 404)
        self.assertEqual(self.request(self.interviewer, path, "POST", {})[0], 403)
        self.assertEqual(self.request(self.admin, path, "POST", {})[0], 409)
        self.request(self.admin, f"/studies/{sid}/questionnaire", "POST", {"questions": [
            {"type": "single", "label": "Редкая категория", "options": ["Часто", "Редко"]},
            {"type": "single", "label": "Безопасный вопрос", "options": ["Да", "Нет"]},
            {"type": "text", "label": "Открытый ответ"},
            {"type": "number", "label": "Оценка"},
        ]})
        for index in range(10):
            self.assertEqual(self.request(self.admin, f"/studies/{sid}/responses", "POST", {"answers": {
                "q1": "Редко" if index < 2 else "Часто", "q2": "Да" if index < 5 else "Нет",
                "q3": "секрет респондента", "q4": "7"}})[0], 201)
        with patch.dict(os.environ, {"DEEPSEEK_API_KEY": "test-placeholder"}), patch("server.urlopen") as urlopen:
            urlopen.return_value.__enter__.return_value.read.return_value = json.dumps({
                "choices": [{"message": {"content": "Резюме\n\nОписание итогов.\n\nРекомендации\n\nПроверить дополнительно."}}
            ]}).encode()
            status, result = self.request(self.admin, path, "POST", {})
            self.assertEqual(status, 200)
            self.assertFalse(result["report"]["stale"])
            request = urlopen.call_args.args[0]
            data = json.loads(request.data)
            self.assertEqual(data["model"], "deepseek-flash")
            self.assertEqual(data["thinking"]["type"], "disabled")
            sent = data["messages"][1]["content"]
            self.assertIn("Безопасный вопрос", sent)
            self.assertIn("Оценка", sent)
            self.assertIn("Изучить проблему", sent)
            self.assertIn("Оценить последствия", sent)
            self.assertNotIn("секрет респондента", sent)
            self.assertIn("Открытый ответ", sent)
            self.assertIn("Редкая категория", sent)
            self.assertNotIn("Редко", sent)
            self.assertEqual(len(result["report"]["snapshot"]["questions"]), 4)
            self.assertNotIn("rows", result["report"]["snapshot"]["questions"][0])
            self.assertIn("note", result["report"]["snapshot"]["questions"][2])
            self.assertEqual(self.request(self.interviewer, path)[0], 403)
            self.assertEqual(self.request(self.interviewer, path + ".pdf", binary=True)[0], 403)
            pdf = self.request(self.admin, path + ".pdf", binary=True)[1]
            self.assertTrue(pdf.startswith(b"%PDF-"))
            self.assertIn(b"TimesNewRoman", pdf)
        self.assertEqual(self.request(self.admin, path)[1]["report"]["content"].splitlines()[0], "Резюме")
        self.assertEqual(self.request(self.admin, f"/studies/{sid}/report")[1]["goal"], "Изучить проблему")
        self.assertEqual(self.request(self.admin, f"/studies/{sid}", "PATCH", {
            "title": "Проверка", "description": "", "goal": "Уточнить цель",
            "tasks": "Измерить результаты", "due_date": "", "responsible_id": None,
            "stage": "development"})[0], 200)
        self.assertTrue(self.request(self.admin, path)[1]["report"]["stale"])
        self.assertEqual(self.request(self.admin, path)[1]["report"]["snapshot"]["goal"], "Изучить проблему")
        self.request(self.admin, f"/studies/{sid}/responses", "POST", {"answers": {"q1": "Часто", "q2": "Да", "q3": "новый текст", "q4": "7"}})
        self.assertTrue(self.request(self.admin, path)[1]["report"]["stale"])
        with patch("server.analytical_pdf_bytes", wraps=analytical_pdf_bytes) as pdf:
            self.assertEqual(self.request(self.admin, path + ".pdf", binary=True)[0], 200)
            self.assertTrue(pdf.call_args.args[0]["stale"])

    def test_analytical_report_without_eligible_questions(self):
        self.request(self.admin, "/setup", "POST", {"name": "Админ", "email": "admin@test.org", "password": "secure-pass-123"})
        sid = self.request(self.admin, "/studies", "POST", {"title": "Открытые ответы", "description": "", "due_date": "", "responsible_id": None})[1]["id"]
        self.request(self.admin, f"/studies/{sid}/questionnaire", "POST", {"questions": [
            {"type": "text", "label": "Комментарий"}]})
        for _ in range(10):
            self.assertEqual(self.request(self.admin, f"/studies/{sid}/responses", "POST", {
                "answers": {"q1": "конфиденциальный ответ"}})[0], 201)
        path = f"/studies/{sid}/analytical-report"
        with patch.dict(os.environ, {"DEEPSEEK_API_KEY": "test-placeholder"}), patch("server.urlopen") as urlopen:
            urlopen.return_value.__enter__.return_value.read.return_value = json.dumps({
                "choices": [{"message": {"content": "Нет данных по вопросам. Доступен только общий объём анкет."}}
            ]}).encode()
            status, result = self.request(self.admin, path, "POST", {})
            self.assertEqual(status, 200)
            self.assertEqual(len(result["report"]["snapshot"]["questions"]), 1)
            sent = json.loads(urlopen.call_args.args[0].data)["messages"][1]["content"]
            self.assertNotIn("конфиденциальный ответ", sent)
            self.assertIn("Комментарий", sent)
            self.assertTrue(self.request(self.admin, path + ".pdf", binary=True)[1].startswith(b"%PDF-"))

    def test_presentation_generation_edit_export_and_staleness(self):
        self.request(self.admin, "/setup", "POST", {"name": "Админ", "email": "admin@test.org", "password": "secure-pass-123"})
        self.request(self.admin, "/users", "POST", {"name": "Интервьюер", "email": "field@test.org", "password": "secure-pass-456", "role": "interviewer"})
        self.request(self.interviewer, "/login", "POST", {"email": "field@test.org", "password": "secure-pass-456"})
        sid = self.request(self.admin, "/studies", "POST", {"title": "Исследование"})[1]["id"]
        path = f"/studies/{sid}/presentation"
        self.assertEqual(self.request(self.admin, path)[1]["presentation"], None)
        self.assertEqual(self.request(self.interviewer, path)[0], 403)
        self.assertEqual(self.request(self.interviewer, path, "POST", {"slide_count": 3})[0], 403)
        self.assertEqual(self.request(self.admin, path + ".pptx", binary=True)[0], 404)
        colors = {"background": "#ffffff", "text": "#193a54", "accent": "#14527c"}
        request = {"slide_count": 3, "prompt": "Сосредоточься на ограничениях", "colors": colors}
        self.assertEqual(self.request(self.admin, path, "POST", request)[0], 409)
        self.request(self.admin, f"/studies/{sid}/questionnaire", "POST", {"questions": [
            {"type": "text", "label": "Комментарий"}]})
        for _ in range(10):
            self.request(self.admin, f"/studies/{sid}/responses", "POST", {"answers": {"q1": "личный ответ"}})
        with patch("server.deepseek_key", return_value="placeholder"), patch("server.deepseek_answer", return_value="Обзор и выводы"):
            self.assertEqual(self.request(self.admin, f"/studies/{sid}/analytical-report", "POST", {})[0], 200)
        outline = json.dumps({"slides": [{"title": f"Слайд {i}", "bullets": ["Развернутый вывод по отчёту с объяснением результатов."],
                                          "summary": "Основной вывод по выборке"} for i in range(1, 4)]})
        with patch("server.deepseek_key", return_value="placeholder"), patch("server.deepseek_answer", return_value=outline) as ai:
            status, response = self.request(self.admin, path, "POST", request)
            self.assertEqual(status, 200)
            sent = json.loads(ai.call_args.args[0]["messages"][1]["content"])
            self.assertEqual(sent["wishes"], request["prompt"])
            self.assertNotIn("личный ответ", json.dumps(sent, ensure_ascii=False))
        deck = response["presentation"]
        self.assertEqual(len(deck["slides"]), 3)
        self.assertEqual(deck["slides"][0]["design"]["variant"], "cover")
        self.assertEqual(deck["slides"][1]["design"]["variant"], "content")
        self.assertEqual(deck["slides"][2]["design"]["variant"], "closing")
        self.assertEqual(deck["slides"][1]["elements"][3]["text"], "Основной вывод по выборке")
        self.assertTrue(deck["slides"][0]["background"].startswith("#"))
        self.assertFalse(self.request(self.admin, path)[1]["stale"])
        deck["slides"][0]["elements"][0].update(text="Перемещённый заголовок", x=10, y=15)
        self.assertEqual(self.request(self.admin, path, "PUT", deck)[0], 200)
        self.assertEqual(self.request(self.admin, path)[1]["presentation"]["slides"][0]["elements"][0]["x"], 10)
        pptx = self.request(self.admin, path + ".pptx", binary=True)[1]
        slides = Presentation(io.BytesIO(pptx)).slides
        self.assertEqual(len(slides), 3)
        self.assertIn("Перемещённый заголовок", [shape.text for shape in slides[0].shapes if shape.has_text_frame])
        self.assertIn("Основной вывод по выборке", [shape.text for shape in slides[1].shapes if shape.has_text_frame])
        self.assertGreater(len(slides[1].shapes), len(deck["slides"][1]["elements"]))
        self.assertFalse(slides[1].shapes[0].has_text_frame and slides[1].shapes[0].text)
        self.assertEqual(slides[1].shapes[0].fill.fore_color.rgb,
                         RGBColor.from_string(deck["slides"][1]["design"]["panel"][1:]))
        deck["slides"][1]["elements"][3]["text"] = "Исправленный вывод"
        deck["slides"][1]["elements"][3]["style"]["size"] = 18
        deck["slides"][1]["background"] = "#eaf4f8"
        self.assertEqual(self.request(self.admin, path, "PUT", deck)[0], 200)
        edited = Presentation(io.BytesIO(self.request(self.admin, path + ".pptx", binary=True)[1])).slides[1]
        self.assertEqual(edited.background.fill.fore_color.rgb, RGBColor(234, 244, 248))
        self.assertEqual(next(shape for shape in edited.shapes if shape.has_text_frame and shape.text == "Исправленный вывод")
                         .text_frame.paragraphs[0].runs[0].font.size.pt, 18)
        self.assertEqual(self.request(self.admin, path, "PUT", {**deck, "slides": [
            {**deck["slides"][0], "design": {**deck["slides"][0]["design"], "panel": "invalid"}},
            *deck["slides"][1:]]})[0], 400)
        self.assertEqual(self.request(self.admin, path, "PUT", {**deck, "colors": {**colors, "accent": "red"}})[0], 400)
        self.assertEqual(self.request(self.admin, path, "PUT", {**deck, "slides": [{"elements": [{"kind": "body", "text": "x", "x": -1, "y": 0, "w": 20, "h": 10}]}] * 3})[0], 400)
        with patch("server.deepseek_key", return_value="placeholder"), patch("server.deepseek_answer", return_value="Новый текст записки"):
            self.assertEqual(self.request(self.admin, f"/studies/{sid}/analytical-report", "POST", {})[0], 200)
        self.assertTrue(self.request(self.admin, path)[1]["stale"])

    def test_pptx_template_applies_fonts_background_and_layout_without_copying_text(self):
        self.request(self.admin, "/setup", "POST", {"name": "Админ", "email": "admin@test.org", "password": "secure-pass-123"})
        self.request(self.admin, "/users", "POST", {"name": "Интервьюер", "email": "field@test.org", "password": "secure-pass-456", "role": "interviewer"})
        self.request(self.interviewer, "/login", "POST", {"email": "field@test.org", "password": "secure-pass-456"})
        sid = self.request(self.admin, "/studies", "POST", {"title": "Образец"})[1]["id"]
        path = f"/studies/{sid}/presentation-template"
        self.assertEqual(self.request(self.admin, path)[1]["template"], None)
        self.assertEqual(self.request(self.interviewer, path)[0], 403)
        self.assertEqual(self.request(self.admin, path, "POST", {"name": "ref.pptx", "content": "@@@"})[0], 400)
        self.assertEqual(self.request(self.admin, path, "POST", {"name": "ref.pptx", "content": base64.b64encode(b"broken").decode()})[0], 400)
        sample = Presentation()
        sample.slide_width, sample.slide_height = Inches(13.333), Inches(7.5)
        cover = sample.slides.add_slide(sample.slide_layouts[6])
        cover.background.fill.solid()
        cover.background.fill.fore_color.rgb = RGBColor(32, 32, 64)
        shape = cover.shapes.add_textbox(Inches(2), Inches(1), Inches(9), Inches(1.5))
        run = shape.text_frame.paragraphs[0].add_run()
        run.text = "Не копировать содержимое образца"
        run.font.name, run.font.size = "Georgia", Pt(34)
        run.font.color.rgb = RGBColor(255, 204, 0)
        content_slide = sample.slides.add_slide(sample.slide_layouts[6])
        content_slide.background.fill.solid()
        content_slide.background.fill.fore_color.rgb = RGBColor(242, 243, 244)
        heading = content_slide.shapes.add_textbox(Inches(1.4), Inches(.7), Inches(10), Inches(1))
        heading_run = heading.text_frame.paragraphs[0].add_run()
        heading_run.text = "Заголовок образца"
        heading_run.font.name, heading_run.font.size = "Georgia", Pt(31)
        heading_run.font.bold = True
        heading_run.font.color.rgb = RGBColor(96, 32, 64)
        body = content_slide.shapes.add_textbox(Inches(1.5), Inches(2), Inches(9), Inches(4))
        body_run = body.text_frame.paragraphs[0].add_run()
        body_run.text = "Приватный текст образца"
        body_run.font.name, body_run.font.size = "Times New Roman", Pt(19)
        body_run.font.color.rgb = RGBColor(32, 64, 96)
        buffer = io.BytesIO()
        sample.save(buffer)
        upload = {"name": "style.pptx", "content": base64.b64encode(buffer.getvalue()).decode()}
        self.assertEqual(self.request(self.interviewer, path, "POST", upload)[0], 403)
        self.assertEqual(self.request(self.admin, path, "POST", upload)[0], 200)
        self.assertEqual(self.request(self.admin, path)[1]["template"]["colors"]["background"], "#f2f3f4")
        with closing(sqlite3.connect(self.db_path)) as db:
            style = json.loads(db.execute("SELECT style FROM presentation_templates WHERE study_id = ?", (sid,)).fetchone()[0])
        self.assertNotIn("Приватный текст", json.dumps(style, ensure_ascii=False))
        self.assertEqual(style["content"]["body"]["font"], "Times New Roman")
        self.assertEqual(style["cover"]["background"], "#202040")
        self.request(self.admin, f"/studies/{sid}/questionnaire", "POST", {"questions": [{"type": "single", "label": "Выбор", "options": ["Да", "Нет"]}]})
        for _ in range(10):
            self.request(self.admin, f"/studies/{sid}/responses", "POST", {"answers": {"q1": "Да"}})
        with patch("server.deepseek_key", return_value="placeholder"), patch("server.deepseek_answer", return_value="Итоги по выбору"):
            self.assertEqual(self.request(self.admin, f"/studies/{sid}/analytical-report", "POST", {})[0], 200)
        outline = json.dumps({"slides": [{"title": f"Тест {i}", "bullets": ["Итог"]} for i in range(3)]})
        with patch("server.deepseek_key", return_value="placeholder"), patch("server.deepseek_answer", return_value=outline) as ai:
            status, result = self.request(self.admin, f"/studies/{sid}/presentation", "POST", {
                "slide_count": 3, "prompt": "Кратко", "colors": style["colors"]})
            self.assertEqual(status, 200)
            self.assertNotIn("Приватный текст образца", json.dumps(ai.call_args.args[0], ensure_ascii=False))
        deck = result["presentation"]
        self.assertEqual(deck["slides"][0]["elements"][0]["style"]["font"], "Georgia")
        self.assertEqual(deck["slides"][0]["elements"][0]["style"]["color"], "#ffcc00")
        self.assertEqual(deck["slides"][0]["background"], "#202040")
        self.assertEqual(deck["slides"][1]["elements"][1]["style"]["size"], 19)
        self.assertGreater(deck["slides"][0]["elements"][0]["x"], 10)
        slides = Presentation(io.BytesIO(self.request(self.admin, f"/studies/{sid}/presentation.pptx", binary=True)[1])).slides
        self.assertEqual(slides[0].background.fill.fore_color.rgb, RGBColor(32, 32, 64))
        cover_title = next(shape for shape in slides[0].shapes if shape.has_text_frame and shape.text == "Тест 0")
        content_body = next(shape for shape in slides[1].shapes if shape.has_text_frame and "Итог" in shape.text)
        self.assertEqual(cover_title.text_frame.paragraphs[0].runs[0].font.name, "Georgia")
        self.assertEqual(cover_title.text_frame.paragraphs[0].runs[0].font.color.rgb, RGBColor(255, 204, 0))
        self.assertEqual(content_body.text_frame.paragraphs[0].runs[0].font.name, "Times New Roman")
        self.assertNotIn("образца", " ".join(shape.text for slide in slides for shape in slide.shapes if shape.has_text_frame))
        deck["slides"][1]["elements"][1]["text"] = "Изменённый текст"
        self.assertEqual(self.request(self.admin, f"/studies/{sid}/presentation", "PUT", deck)[0], 200)
        saved = self.request(self.admin, f"/studies/{sid}/presentation")[1]["presentation"]
        self.assertEqual(saved["slides"][1]["elements"][1]["style"]["font"], "Times New Roman")
        self.assertEqual(self.request(self.admin, path, "POST", {"name": "bad.pptx", "content": base64.b64encode(b"bad").decode()})[0], 400)
        self.assertEqual(self.request(self.admin, path)[1]["template"]["name"], "style.pptx")
        self.assertEqual(self.request(self.admin, path, "DELETE", {})[0], 200)
        self.assertIsNone(self.request(self.admin, path)[1]["template"])
        self.assertEqual(self.request(self.admin, path, "POST", upload)[0], 200)
        self.assertEqual(self.request(self.admin, f"/studies/{sid}", "DELETE", {})[0], 200)
        with closing(sqlite3.connect(self.db_path)) as db:
            self.assertEqual(db.execute("SELECT count(*) FROM presentation_templates WHERE study_id = ?", (sid,)).fetchone()[0], 0)

    def test_legacy_presentation_still_exports_without_modern_decoration(self):
        deck = {"colors": {"background": "#ffffff", "text": "#193a54", "accent": "#14527c"},
                "slides": [{"elements": [{"kind": "title", "text": "Старый слайд", "x": 6, "y": 8,
                                          "w": 88, "h": 20}]} for _ in range(3)]}
        slides = Presentation(io.BytesIO(pptx_bytes(deck))).slides
        self.assertEqual(len(slides[0].shapes), 1)
        self.assertEqual(slides[0].shapes[0].text, "Старый слайд")


if __name__ == "__main__":
    unittest.main()
