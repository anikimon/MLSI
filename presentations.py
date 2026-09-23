"""Editable slide layout and PowerPoint export for analytical reports."""

import io
import math
import re

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.util import Inches, Pt


def validate_colors(colors):
    if not isinstance(colors, dict) or any(not isinstance(colors.get(key), str) or
                                          not re.fullmatch(r"#[0-9a-fA-F]{6}", colors[key])
                                          for key in ("background", "text", "accent")):
        raise ValueError("Выберите цвета фона, текста и акцента")
    return {key: colors[key].lower() for key in ("background", "text", "accent")}


def validate_deck(deck):
    if not isinstance(deck, dict) or not isinstance(deck.get("slides"), list) or not 3 <= len(deck["slides"]) <= 20:
        raise ValueError("Презентация должна содержать от 3 до 20 слайдов")
    cleaned = {"colors": validate_colors(deck.get("colors")), "slides": []}
    for slide in deck["slides"]:
        if not isinstance(slide, dict) or not isinstance(slide.get("elements"), list) or not 1 <= len(slide["elements"]) <= 8:
            raise ValueError("Каждый слайд должен содержать от 1 до 8 текстовых объектов")
        elements = []
        for element in slide["elements"]:
            if not isinstance(element, dict) or element.get("kind") not in ("title", "body") or not isinstance(element.get("text"), str) or not element["text"].strip() or len(element["text"]) > 1200:
                raise ValueError("Некорректный текст слайда")
            coordinates = [element.get(key) for key in ("x", "y", "w", "h")]
            if any(type(value) not in (int, float) or not math.isfinite(value) or not 0 <= value <= 100 for value in coordinates):
                raise ValueError("Некорректные координаты объекта")
            x, y, w, h = coordinates
            if w < 5 or h < 5 or x + w > 100 or y + h > 100:
                raise ValueError("Объект должен помещаться на слайде")
            elements.append({"kind": element["kind"], "text": element["text"].strip(),
                             "x": x, "y": y, "w": w, "h": h})
        cleaned["slides"].append({"elements": elements})
    return cleaned


def deck_from_outline(outline, count, colors):
    if not isinstance(outline, dict) or not isinstance(outline.get("slides"), list) or len(outline["slides"]) != count:
        raise ValueError("ИИ вернул неверное количество слайдов")
    slides = []
    for item in outline["slides"]:
        if not isinstance(item, dict) or not isinstance(item.get("title"), str) or not isinstance(item.get("bullets"), list):
            raise ValueError("ИИ вернул некорректную структуру слайдов")
        title = item["title"].strip()[:180]
        bullets = item["bullets"]
        if not title or len(bullets) > 6 or any(not isinstance(bullet, str) or not bullet.strip() or len(bullet) > 180 for bullet in bullets):
            raise ValueError("ИИ вернул слишком длинный или пустой текст слайда")
        body = "\n".join("• " + bullet.strip() for bullet in bullets)
        elements = [{"kind": "title", "text": title, "x": 6, "y": 8, "w": 88, "h": 20}]
        if body:
            elements.append({"kind": "body", "text": body, "x": 8, "y": 33, "w": 84, "h": 57})
        slides.append({"elements": elements})
    return validate_deck({"colors": colors, "slides": slides})


def pptx_bytes(deck):
    deck = validate_deck(deck)
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    colors = {key: RGBColor.from_string(value[1:]) for key, value in deck["colors"].items()}
    for item in deck["slides"]:
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = colors["background"]
        for element in item["elements"]:
            shape = slide.shapes.add_textbox(int(presentation.slide_width * element["x"] / 100),
                                             int(presentation.slide_height * element["y"] / 100),
                                             int(presentation.slide_width * element["w"] / 100),
                                             int(presentation.slide_height * element["h"] / 100))
            frame = shape.text_frame
            frame.word_wrap = True
            frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
            frame.margin_left = frame.margin_right = Inches(.06)
            frame.margin_top = frame.margin_bottom = Inches(.04)
            for line_number, line in enumerate(element["text"].splitlines()):
                paragraph = frame.paragraphs[0] if line_number == 0 else frame.add_paragraph()
                paragraph.text = line
                paragraph.alignment = PP_ALIGN.LEFT
                paragraph.space_after = Pt(9)
                for run in paragraph.runs:
                    run.font.name = "Aptos"
                    run.font.size = Pt(28 if element["kind"] == "title" else 17)
                    run.font.bold = element["kind"] == "title"
                    run.font.color.rgb = colors["accent" if element["kind"] == "title" else "text"]
    output = io.BytesIO()
    presentation.save(output)
    return output.getvalue()
